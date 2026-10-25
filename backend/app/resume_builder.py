"""ATS-friendly resume (.docx) built from a seeker profile.

ATS rules followed: one column; no tables, text boxes, images, headers or footers; standard section names
(Summary, Skills, Professional Experience, Projects, Education, Certifications, Languages); plain bullets;
a standard font; contact details and links as plain text at the top.
python-docx is imported inside build_docx() so the API still starts if the package is not installed yet.
"""
import re
from io import BytesIO

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
_BAD = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")        # characters Word documents cannot hold


def _clean(v) -> str:
    return _BAD.sub("", str(v or "")).strip()


def _month(v) -> str:
    v = _clean(v)
    m = re.fullmatch(r"(\d{4})-(\d{1,2})", v)
    if m and 1 <= int(m.group(2)) <= 12:
        return f"{MONTHS[int(m.group(2)) - 1]} {m.group(1)}"
    return v


def _bullets(text) -> list[str]:
    lines = [re.sub(r"^[\s•\-\*·▪–]+", "", ln).strip() for ln in _clean(text).splitlines()]
    lines = [ln for ln in lines if ln]
    if len(lines) == 1 and len(lines[0]) > 160:                 # one long paragraph: one bullet per sentence
        lines = [s.strip() for s in re.split(r"(?<=[.!?])\s+", lines[0]) if s.strip()]
    return lines[:8]


def build_docx(p: dict) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    NAVY = RGBColor(0x1F, 0x3A, 0x6B)
    name = _clean(p.get("full_name")) or _clean(p.get("email")).split("@")[0] or "Candidate"

    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(0.8)
    sec.top_margin = sec.bottom_margin = Inches(0.7)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(2)
    doc.core_properties.title = f"{name} - Resume"
    doc.core_properties.author = name

    def line(text, bold=False, italic=False, size=None, align=None, after=None, color=None):
        para = doc.add_paragraph()
        run = para.add_run(_clean(text))
        run.bold, run.italic = bold, italic
        if size:
            run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = color
        if align is not None:
            para.alignment = align
        if after is not None:
            para.paragraph_format.space_after = Pt(after)
        return para

    def heading(text):
        para = doc.add_heading(level=2)
        run = para.add_run(text.upper())
        run.font.name, run.font.size, run.bold, run.font.color.rgb = "Calibri", Pt(11.5), True, NAVY
        pPr = para._p.get_or_add_pPr()                         # thin rule under the heading (a border, not a table)
        border, bottom = OxmlElement("w:pBdr"), OxmlElement("w:bottom")
        for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "1F3A6B")):
            bottom.set(qn(k), v)
        border.append(bottom)
        pPr.append(border)
        para.paragraph_format.space_before = Pt(10)
        para.paragraph_format.space_after = Pt(3)

    def bullet(text):
        para = doc.add_paragraph(style="List Bullet")
        para.add_run(_clean(text))
        para.paragraph_format.space_after = Pt(1)

    # header: name, target title, contact line, links line
    line(name, bold=True, size=20, align=WD_ALIGN_PARAGRAPH.CENTER, after=0, color=NAVY)
    titles = [t for t in (p.get("job_titles") or []) if _clean(t)]
    if titles:
        line(_clean(titles[0]), size=12, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    contact = [_clean(p.get(k)) for k in ("email", "phone", "location") if _clean(p.get(k))]
    if contact:
        line(" | ".join(contact), align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    links = [f"{lbl}: {_clean(p.get(k))}" for lbl, k in (("LinkedIn", "linkedin_url"), ("GitHub", "github_url"), ("Portfolio", "portfolio_url")) if _clean(p.get(k))]
    if links:
        line(" | ".join(links), align=WD_ALIGN_PARAGRAPH.CENTER, after=0)

    if _clean(p.get("summary")):
        heading("Summary")
        line(p["summary"])

    skills = [_clean(s) for s in (p.get("skills") or []) if _clean(s)]
    if skills:
        heading("Skills")
        line(", ".join(skills))

    exp = [e for e in (p.get("experience") or []) if _clean(e.get("title"))]
    if exp:
        heading("Professional Experience")
        for e in exp:
            head = _clean(e.get("title")) + (f" | {_clean(e.get('company'))}" if _clean(e.get("company")) else "")
            line(head, bold=True, after=0)
            start, end = _month(e.get("start")), "Present" if e.get("current") else _month(e.get("end"))
            if start or end:
                line(f"{start} - {end}".strip(" -"), italic=True, after=1)
            for b in _bullets(e.get("description")):
                bullet(b)

    projects = [x for x in (p.get("projects") or []) if _clean(x.get("name"))]
    if projects:
        heading("Projects")
        for x in projects:
            stack = ", ".join(_clean(t) for t in (x.get("tech_stack") or []) if _clean(t))
            line(_clean(x.get("name")) + (f" | {stack}" if stack else ""), bold=True, after=0)
            if _clean(x.get("url")):
                line(_clean(x.get("url")), after=1)
            for b in _bullets(x.get("description")):
                bullet(b)

    edu = [e for e in (p.get("education") or []) if _clean(e.get("degree"))]
    if edu:
        heading("Education")
        for e in edu:
            parts = [_clean(e.get("degree"))] + [_clean(e.get(k)) for k in ("institution", "year") if _clean(e.get(k))]
            line(" | ".join(parts))

    certs = [c for c in (p.get("certifications") or []) if _clean(c.get("name"))]
    if certs:
        heading("Certifications")
        for c in certs:
            parts = [_clean(c.get("name"))] + [_clean(c.get(k)) for k in ("issuer", "year") if _clean(c.get(k))]
            line(" | ".join(parts))

    langs = [x for x in (p.get("languages") or []) if _clean(x.get("language"))]
    if langs:
        heading("Languages")
        line(", ".join(_clean(x.get("language")) + (f" ({_clean(x.get('level'))})" if _clean(x.get("level")) else "") for x in langs))

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
